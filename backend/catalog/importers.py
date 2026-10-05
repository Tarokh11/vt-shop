from dataclasses import dataclass
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .models import (
    AttributeDefinition,
    AttributeValue,
    Brand,
    Category,
    CategoryAttributeDefinition,
    Collection,
    CollectionProduct,
    InventoryAdjustment,
    Product,
    ProductAttributeValue,
    ProductOptionDefinition,
    ProductVariant,
    VariantOptionValue,
)

HEADERS = (
    "product_slug",
    "product_name",
    "description",
    "category_slugs",
    "brand_slug",
    "collection_slugs",
    "is_published",
    "sku",
    "variant_name",
    "price_irr",
    "stock_quantity",
    "is_active",
    "is_default",
    "attributes",
    "options",
)
REQUIRED_HEADERS = {"product_slug", "product_name", "category_slugs", "sku", "price_irr"}


@dataclass(frozen=True)
class CatalogRow:
    number: int
    product_slug: str
    product_name: str
    description: str
    category_slugs: tuple[str, ...]
    brand_slug: str
    collection_slugs: tuple[str, ...]
    is_published: bool
    sku: str
    variant_name: str
    price_irr: int
    stock_quantity: int
    is_active: bool
    is_default: bool
    attributes: dict[str, str]
    options: dict[str, str]


@dataclass(frozen=True)
class ImportResult:
    products_created: int
    products_updated: int
    variants_created: int
    variants_updated: int
    stock_adjustments: int


class CatalogImportError(Exception):
    def __init__(self, errors):
        self.errors = errors
        super().__init__("Catalog workbook contains invalid data.")


def _text(value):
    return "" if value is None else str(value).strip()


def _list(value):
    return tuple(item.strip() for item in _text(value).split(",") if item.strip())


def _mapping(value):
    result = {}
    for item in _text(value).split("|"):
        item = item.strip()
        if not item:
            continue
        if ":" not in item:
            raise ValueError(f"'{item}' باید به شکل definition:value باشد.")
        definition, option = (part.strip() for part in item.split(":", 1))
        if not definition or not option:
            raise ValueError(f"'{item}' ناقص است.")
        result[definition] = option
    return result


def _boolean(value, default):
    if value in (None, ""):
        return default
    normalized = _text(value).lower()
    if normalized in {"1", "true", "yes", "y", "بله"}:
        return True
    if normalized in {"0", "false", "no", "n", "خیر"}:
        return False
    raise ValueError(f"مقدار بولی '{value}' معتبر نیست.")


def _integer(value, field, minimum=0):
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field} باید عدد صحیح باشد.") from exc
    if number < minimum:
        raise ValueError(f"{field} نمی‌تواند کمتر از {minimum} باشد.")
    return number


def _read_rows(file):
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:
        raise CatalogImportError(["فایل اکسل خوانده نشد یا خراب است."]) from exc
    sheet = workbook["Products"] if "Products" in workbook.sheetnames else workbook.active
    values = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(values)
    except StopIteration as exc:
        raise CatalogImportError(["فایل اکسل خالی است."]) from exc
    headers = [_text(value) for value in raw_headers]
    missing = REQUIRED_HEADERS - set(headers)
    if missing:
        raise CatalogImportError([f"ستون‌های ضروری وجود ندارند: {', '.join(sorted(missing))}"])

    rows, errors = [], []
    for number, values_row in enumerate(values, start=2):
        data = dict(zip(headers, values_row, strict=False))
        if not any(value not in (None, "") for value in values_row):
            continue
        try:
            product_slug = _text(data.get("product_slug"))
            product_name = _text(data.get("product_name"))
            categories = _list(data.get("category_slugs"))
            sku = _text(data.get("sku"))
            if not product_slug or not product_name or not categories or not sku:
                raise ValueError("شناسه محصول، نام محصول، دسته‌بندی و SKU ضروری هستند.")
            rows.append(
                CatalogRow(
                    number=number,
                    product_slug=product_slug,
                    product_name=product_name,
                    description=_text(data.get("description")),
                    category_slugs=categories,
                    brand_slug=_text(data.get("brand_slug")),
                    collection_slugs=_list(data.get("collection_slugs")),
                    is_published=_boolean(data.get("is_published"), False),
                    sku=sku,
                    variant_name=_text(data.get("variant_name")),
                    price_irr=_integer(data.get("price_irr"), "price_irr", 1),
                    stock_quantity=_integer(data.get("stock_quantity") or 0, "stock_quantity"),
                    is_active=_boolean(data.get("is_active"), True),
                    is_default=_boolean(data.get("is_default"), False),
                    attributes=_mapping(data.get("attributes")),
                    options=_mapping(data.get("options")),
                )
            )
        except ValueError as exc:
            errors.append(f"ردیف {number}: {exc}")
    if not rows and not errors:
        errors.append("فایل هیچ ردیف محصولی ندارد.")
    if errors:
        raise CatalogImportError(errors)
    return rows


def _validate_references(rows):
    errors = []
    category_slugs = {slug for row in rows for slug in row.category_slugs}
    brand_slugs = {row.brand_slug for row in rows if row.brand_slug}
    collection_slugs = {slug for row in rows for slug in row.collection_slugs}
    definition_slugs = {
        slug for row in rows for slug in (*row.attributes.keys(), *row.options.keys())
    }
    categories = {item.slug: item for item in Category.objects.filter(slug__in=category_slugs)}
    brands = {item.slug: item for item in Brand.objects.filter(slug__in=brand_slugs)}
    collections = {
        item.slug: item for item in Collection.objects.filter(slug__in=collection_slugs)
    }
    definitions = {
        item.slug: item for item in AttributeDefinition.objects.filter(slug__in=definition_slugs)
    }
    value_pairs = {
        (item.definition.slug, item.slug): item
        for item in AttributeValue.objects.filter(
            definition__slug__in=definition_slugs, is_active=True
        ).select_related("definition")
    }

    seen_skus = set()
    defaults = {}
    product_metadata = {}
    for row in rows:
        if row.sku in seen_skus:
            errors.append(f"ردیف {row.number}: SKU تکراری '{row.sku}' در فایل.")
        seen_skus.add(row.sku)
        existing_variant = (
            ProductVariant.objects.filter(sku=row.sku).select_related("product").first()
        )
        if existing_variant and existing_variant.product.slug != row.product_slug:
            errors.append(
                f"ردیف {row.number}: SKU '{row.sku}' متعلق به محصول دیگری است."
            )
        defaults[row.product_slug] = defaults.get(row.product_slug, 0) + int(row.is_default)
        signature = (
            row.product_name,
            row.description,
            row.category_slugs,
            row.brand_slug,
            row.collection_slugs,
            row.is_published,
            row.attributes,
        )
        previous = product_metadata.setdefault(row.product_slug, signature)
        if previous != signature:
            errors.append(f"ردیف {row.number}: اطلاعات تکرارشده محصول یکسان نیست.")
        for slug in row.category_slugs:
            if slug not in categories:
                errors.append(f"ردیف {row.number}: دسته‌بندی '{slug}' وجود ندارد.")
        if row.brand_slug and row.brand_slug not in brands:
            errors.append(f"ردیف {row.number}: برند '{row.brand_slug}' وجود ندارد.")
        for slug in row.collection_slugs:
            if slug not in collections:
                errors.append(f"ردیف {row.number}: مجموعه '{slug}' وجود ندارد.")
        for definition_slug, value_slug in {**row.attributes, **row.options}.items():
            if definition_slug not in definitions:
                errors.append(f"ردیف {row.number}: ویژگی '{definition_slug}' وجود ندارد.")
            elif (definition_slug, value_slug) not in value_pairs:
                errors.append(
                    f"ردیف {row.number}: مقدار '{definition_slug}:{value_slug}' وجود ندارد."
                )
            elif not CategoryAttributeDefinition.objects.filter(
                category__slug__in=row.category_slugs,
                definition__slug=definition_slug,
            ).exists():
                errors.append(
                    f"ردیف {row.number}: ویژگی '{definition_slug}' برای دسته‌بندی تنظیم نشده."
                )
        overlap = row.attributes.keys() & row.options.keys()
        if overlap:
            errors.append(
                f"ردیف {row.number}: ویژگی نمی‌تواند هم ثابت و هم گزینه تنوع باشد: "
                f"{', '.join(sorted(overlap))}"
            )
    for slug, count in defaults.items():
        if count > 1:
            errors.append(f"محصول '{slug}' بیش از یک تنوع پیش‌فرض دارد.")
    if errors:
        raise CatalogImportError(errors)
    return categories, brands, collections, definitions, value_pairs


def import_catalog_workbook(file, user):
    rows = _read_rows(file)
    categories, brands, collections, definitions, values = _validate_references(rows)
    counters = {"pc": 0, "pu": 0, "vc": 0, "vu": 0, "stock": 0}
    touched_products = {}

    with transaction.atomic():
        for row in rows:
            product, created = Product.objects.update_or_create(
                slug=row.product_slug,
                defaults={
                    "name": row.product_name,
                    "description": row.description,
                    "brand": brands.get(row.brand_slug),
                    "is_published": False,
                    "creation_source": Product.CreationSource.INVENTORY,
                    "is_archived": False,
                },
            )
            counters["pc" if created else "pu"] += int(row.product_slug not in touched_products)
            touched_products[row.product_slug] = (product, row.is_published)
            product.categories.set(categories[slug] for slug in row.category_slugs)
            for definition_slug, value_slug in row.attributes.items():
                ProductAttributeValue.objects.get_or_create(
                    product=product, value=values[(definition_slug, value_slug)]
                )
            for position, definition_slug in enumerate(row.options, start=1):
                ProductOptionDefinition.objects.update_or_create(
                    product=product,
                    definition=definitions[definition_slug],
                    defaults={"position": position},
                )
            for position, collection_slug in enumerate(row.collection_slugs, start=1):
                CollectionProduct.objects.update_or_create(
                    collection=collections[collection_slug],
                    product=product,
                    defaults={"position": position},
                )
            if row.is_default:
                product.variants.filter(is_default=True).exclude(sku=row.sku).update(is_default=False)
            variant, variant_created = ProductVariant.objects.update_or_create(
                sku=row.sku,
                defaults={
                    "product": product,
                    "name": row.variant_name,
                    "price_irr": row.price_irr,
                    "is_active": row.is_active,
                    "is_default": row.is_default,
                },
            )
            counters["vc" if variant_created else "vu"] += 1
            for definition_slug, value_slug in row.options.items():
                option = ProductOptionDefinition.objects.get(
                    product=product, definition=definitions[definition_slug]
                )
                VariantOptionValue.objects.update_or_create(
                    variant=variant,
                    option=option,
                    defaults={"value": values[(definition_slug, value_slug)]},
                )
            delta = row.stock_quantity - variant.stock_quantity
            if delta:
                InventoryAdjustment.objects.create(
                    variant=variant,
                    quantity_delta=delta,
                    reason=f"Excel catalog import: {row.sku}",
                    created_by=user,
                )
                counters["stock"] += 1

        for product, publish in touched_products.values():
            if not product.variants.filter(is_active=True, is_default=True).exists():
                first_active = product.variants.filter(is_active=True).first()
                if publish and first_active is None:
                    raise ValidationError(f"محصول '{product.slug}' تنوع فعال ندارد.")
                if first_active:
                    first_active.is_default = True
                    first_active.save(update_fields=("is_default", "updated_at"))
            product.is_published = publish
            product.full_clean()
            product.save(update_fields=("is_published", "updated_at"))

    return ImportResult(
        products_created=counters["pc"],
        products_updated=counters["pu"],
        variants_created=counters["vc"],
        variants_updated=counters["vu"],
        stock_adjustments=counters["stock"],
    )


def build_catalog_template():
    workbook = Workbook()
    products = workbook.active
    products.title = "Products"
    products.append(HEADERS)
    products.append(
        (
            "sample-notebook",
            "دفتر نمونه",
            "توضیح کوتاه محصول",
            "notebooks-paper",
            "rooyesh",
            "back-to-school",
            True,
            "SAMPLE-A5-BLUE",
            "آبی / A5",
            390000,
            10,
            True,
            True,
            "ruling:lined|page-count:80",
            "color:blue|paper-size:a5",
        )
    )
    products.freeze_panes = "A2"
    products.auto_filter.ref = "A1:O2"
    for cell in products[1]:
        cell.font = Font(bold=True)
    for column in products.columns:
        products.column_dimensions[column[0].column_letter].width = min(
            max(len(_text(cell.value)) for cell in column) + 3, 34
        )

    guide = workbook.create_sheet("Guide")
    guide.append(("ستون", "راهنما"))
    guide_rows = (
        ("هر ردیف", "یک SKU؛ اطلاعات محصول برای همه تنوع‌های همان محصول تکرار شود."),
        ("category_slugs", "شناسه دسته‌ها با ویرگول جدا شود."),
        ("collection_slugs", "شناسه مجموعه‌ها با ویرگول جدا شود؛ اختیاری."),
        ("attributes", "ویژگی ثابت: definition:value و جداسازی با |"),
        ("options", "گزینه تنوع: definition:value و جداسازی با |"),
        ("stock_quantity", "موجودی هدف؛ اختلاف به‌صورت سند تعدیل موجودی ثبت می‌شود."),
        ("boolean", "true/false یا بله/خیر"),
    )
    for row in guide_rows:
        guide.append(row)
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 82

    reference = workbook.create_sheet("References")
    reference.append(("نوع", "slug", "نام"))
    for item in Category.objects.order_by("position", "name"):
        reference.append(("category", item.slug, item.name))
    for item in Brand.objects.order_by("position", "name"):
        reference.append(("brand", item.slug, item.name))
    for item in Collection.objects.order_by("position", "name"):
        reference.append(("collection", item.slug, item.name))
    for item in AttributeValue.objects.select_related("definition").order_by(
        "definition__position", "position"
    ):
        reference.append(("attribute", f"{item.definition.slug}:{item.slug}", item.label))
    reference.column_dimensions["A"].width = 18
    reference.column_dimensions["B"].width = 38
    reference.column_dimensions["C"].width = 28

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
