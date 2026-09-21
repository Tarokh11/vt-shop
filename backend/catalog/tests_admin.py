from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from accounts.models import User

from .importers import HEADERS, CatalogImportError, build_catalog_template, import_catalog_workbook
from .models import (
    AttributeDefinition,
    AttributeValue,
    Brand,
    Category,
    CategoryAttributeDefinition,
    Collection,
    InventoryAdjustment,
    Product,
    ProductVariant,
)


def workbook_file(rows, name="catalog.xlsx"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    content = BytesIO()
    workbook.save(content)
    return SimpleUploadedFile(
        name,
        content.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class CatalogExcelImportTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_superuser(
            username="admin@example.com", email="admin@example.com", password="test"
        )
        self.category = Category.objects.create(name="دفتر", slug="notebooks-paper")
        self.brand = Brand.objects.create(name="رویش", slug="rooyesh")
        self.collection = Collection.objects.create(name="مدرسه", slug="back-to-school")
        self.material = AttributeDefinition.objects.create(name="جنس", slug="material")
        self.color = AttributeDefinition.objects.create(name="رنگ", slug="color")
        self.paper = AttributeValue.objects.create(
            definition=self.material, label="کاغذ", slug="paper"
        )
        self.blue = AttributeValue.objects.create(
            definition=self.color, label="آبی", slug="blue"
        )
        self.red = AttributeValue.objects.create(
            definition=self.color, label="قرمز", slug="red"
        )
        for definition in (self.material, self.color):
            CategoryAttributeDefinition.objects.create(
                category=self.category, definition=definition
            )

    def rows(self, second_stock=4):
        product = (
            "daily-notebook",
            "دفتر روزانه",
            "دفتر کاربردی برای یادداشت‌های روزانه",
            "notebooks-paper",
            "rooyesh",
            "back-to-school",
            True,
        )
        return [
            (*product, "NOTE-BLUE", "آبی", 390_000, 7, True, True, "material:paper", "color:blue"),
            (
                *product,
                "NOTE-RED",
                "قرمز",
                390_000,
                second_stock,
                True,
                False,
                "material:paper",
                "color:red",
            ),
        ]

    def test_import_creates_and_updates_normalized_catalog_with_audited_stock(self):
        result = import_catalog_workbook(workbook_file(self.rows()), self.staff)

        product = Product.objects.get(slug="daily-notebook")
        self.assertTrue(product.is_published)
        self.assertEqual(product.categories.get(), self.category)
        self.assertEqual(product.brand, self.brand)
        self.assertEqual(product.collections.get(), self.collection)
        self.assertEqual(result.products_created, 1)
        self.assertEqual(result.variants_created, 2)
        self.assertEqual(result.stock_adjustments, 2)
        self.assertEqual(ProductVariant.objects.get(sku="NOTE-BLUE").stock_quantity, 7)
        self.assertEqual(InventoryAdjustment.objects.count(), 2)

        second_result = import_catalog_workbook(
            workbook_file(self.rows(second_stock=6)), self.staff
        )

        self.assertEqual(Product.objects.count(), 1)
        self.assertEqual(ProductVariant.objects.count(), 2)
        self.assertEqual(second_result.products_updated, 1)
        self.assertEqual(second_result.variants_updated, 2)
        self.assertEqual(second_result.stock_adjustments, 1)
        self.assertEqual(ProductVariant.objects.get(sku="NOTE-RED").stock_quantity, 6)
        self.assertEqual(InventoryAdjustment.objects.count(), 3)

    def test_invalid_reference_rolls_back_entire_workbook(self):
        invalid = list(self.rows())
        invalid[1] = (
            "broken-product",
            "محصول نامعتبر",
            "",
            "missing-category",
            "",
            "",
            False,
            "BROKEN-SKU",
            "",
            100,
            1,
            True,
            True,
            "",
            "",
        )

        with self.assertRaises(CatalogImportError):
            import_catalog_workbook(workbook_file(invalid), self.staff)

        self.assertFalse(Product.objects.exists())
        self.assertFalse(ProductVariant.objects.exists())
        self.assertFalse(InventoryAdjustment.objects.exists())

    def test_template_contains_products_guide_and_live_references(self):
        workbook = load_workbook(build_catalog_template(), read_only=True)

        self.assertEqual(workbook.sheetnames, ["Products", "Guide", "References"])
        references = list(workbook["References"].values)
        self.assertIn(("category", "notebooks-paper", "دفتر"), references)
        self.assertIn(("attribute", "color:blue", "آبی"), references)

    def test_admin_exposes_template_and_import_workflow(self):
        self.client.force_login(self.staff)

        changelist = self.client.get(reverse("admin:catalog_product_changelist"))
        template = self.client.get(reverse("admin:catalog_product_import_template"))
        response = self.client.post(
            reverse("admin:catalog_product_import_excel"),
            {"workbook": workbook_file(self.rows())},
        )

        self.assertContains(changelist, "ورود از اکسل")
        self.assertEqual(template.status_code, 200)
        self.assertEqual(
            template["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertRedirects(response, reverse("admin:catalog_product_changelist"))
        self.assertTrue(Product.objects.filter(slug="daily-notebook").exists())
